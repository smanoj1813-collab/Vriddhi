#!/usr/bin/env python3
"""Builds the Vriddhi investor deck from one source of truth.

Outputs (all in marketing/):
  • Vriddhi_Investor_Deck.pptx          — 16:9 PowerPoint you present / send
    (the PDF is built by build_investor_deck_pdf.py from the same content)
  • investor-deck.html                  — same deck in a browser (arrow keys, Ctrl+P → PDF)
  • Vriddhi_Investor_Deck_Speaker_Notes.md — the script, slide by slide

Regenerate:
    python3 -m venv /home/user/.venv
    /home/user/.venv/bin/pip install python-pptx
    /home/user/.venv/bin/python marketing/build_investor_deck.py

Every number in here is traceable to the repository — see docs/INVESTOR/.
Placeholders that only you can fill are written as <angle brackets>; search for "<" before sending.
"""
import html
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
BRAND = os.path.join(ROOT, "brand")
HERO = os.path.join(BRAND, "hero", "vriddhi-mark-hero-light.png")
LOGO_LIGHT = os.path.join(BRAND, "logo", "png", "vriddhi-logo-horizontal-transparent@1600.png")
LOGO_REVERSE = os.path.join(BRAND, "logo", "png", "vriddhi-logo-horizontal-reverse@1600.png")

# ── brand tokens ─────────────────────────────────────────────────────────────
TEAL = "14B8A6"
TEAL_D = "0F766E"
TEAL_3 = "5EEAD4"
TEAL_4 = "2DD4BF"
TEAL_50 = "F0FDFA"
TEAL_100 = "CCFBF1"
SLATE = "0F172A"
SLATE_2 = "1E293B"
SLATE_6 = "475569"
SLATE_5 = "64748B"
SLATE_4 = "94A3B8"
SLATE_3 = "CBD5E1"
BORDER = "E2E8F0"
BG = "F8FAFC"
WHITE = "FFFFFF"
DARK_BG = "0B1220"
FONT = "Inter"

SW, SH = 13.333, 7.5
M = 0.62


# ═════════════════════════════════════════════════════════════════════════════
# 1. CONTENT — the single source of truth for all three outputs
# ═════════════════════════════════════════════════════════════════════════════

from investor_deck_content import SLIDES


# ═════════════════════════════════════════════════════════════════════════════
# 2. PPTX RENDERER
# ═════════════════════════════════════════════════════════════════════════════

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def C(hexc):
    return RGBColor.from_string(hexc)


def slide_bg(color):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C(color)
    return s


def box(slide, x, y, w, h, fill=None, line=None, lw=0.75, round_=None, adj=0.08):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if round_ is not None else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if round_ is not None:
        try:
            shp.adjustments[0] = adj
        except Exception:
            pass
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = C(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = C(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def tb(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = t.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = para.get("align", PP_ALIGN.LEFT)
        if "space_before" in para:
            p.space_before = Pt(para["space_before"])
        if "space_after" in para:
            p.space_after = Pt(para["space_after"])
        p.line_spacing = para.get("line", 1.0)
        for r in para["runs"]:
            run = p.add_run()
            run.text = r["t"]
            f = run.font
            f.name = r.get("font", FONT)
            f.size = Pt(r.get("s", 14))
            f.bold = r.get("b", False)
            f.italic = r.get("i", False)
            f.color.rgb = C(r.get("c", SLATE))
    return t


def P(t, s=14, b=False, i=False, c=SLATE, **kw):
    d = {"t": t, "s": s, "b": b, "i": i, "c": c}
    d.update(kw)
    return d


def header(slide, kicker, title, sub=None, dark=False, title_size=30):
    kc = TEAL_3 if dark else TEAL_D
    tc = WHITE if dark else SLATE
    sc = SLATE_3 if dark else SLATE_5
    box(slide, M, 0.44, 0.5, 0.055, fill=TEAL)
    if kicker:
        tb(slide, M, 0.52, 10, 0.3, [{"runs": [P(kicker, 11, True, False, kc)]}])
    tb(slide, M, 0.82, SW - 2 * M, 1.05, [{"runs": [P(title, title_size, True, False, tc)]}])
    if sub:
        tb(slide, M, 1.52, SW - 2 * M - 0.5, 0.4, [{"runs": [P(sub, 13, False, False, sc)]}])


def bullet_list(slide, x, y, w, items, s=12, gap=7, color=SLATE_6, dot=TEAL, h=4.6, line=1.14):
    paras = []
    for it in items:
        paras.append({
            "runs": [P("●  ", 9, False, False, dot), P(it, s, False, False, color)],
            "space_after": gap,
            "line": line,
        })
    tb(slide, x, y, w, h, paras)


def panel_list(slide, x, y, w, items, s=11.5, gap=9, title_c=SLATE, body_c=SLATE_6, h=4.4):
    """List of (term, detail) pairs."""
    paras = []
    for term, detail in items:
        paras.append({
            "runs": [P(term, s + 0.5, True, False, title_c)],
            "space_after": 1,
            "line": 1.15,
        })
        paras.append({
            "runs": [P(detail, s, False, False, body_c)],
            "space_after": gap,
            "line": 1.2,
        })
    tb(slide, x, y, w, h, paras)


def stat_row(slide, y, stats, box_h=1.5):
    n = len(stats)
    gap = 0.22
    w = (SW - 2 * M - gap * (n - 1)) / n
    for i, (value, label) in enumerate(stats):
        x = M + i * (w + gap)
        box(slide, x, y, w, box_h, fill=WHITE, line=BORDER, round_=True, adj=0.06)
        box(slide, x, y, 0.055, box_h, fill=TEAL)
        # shrink the headline figure when it is long, so the label never collides
        v_size = 23 if len(value) <= 12 else (20 if len(value) <= 18 else 17)
        v_h = 0.5 if v_size >= 23 else 0.62
        tb(slide, x + 0.26, y + 0.16, w - 0.5, v_h,
           [{"runs": [P(value, v_size, True, False, TEAL_D)]}])
        label_y = y + 0.2 + v_h
        tb(slide, x + 0.26, label_y, w - 0.5, max(0.4, y + box_h - label_y - 0.12),
           [{"runs": [P(label, 10.5, False, False, SLATE_6)], "line": 1.18}])


def card_grid(slide, cards, y0=2.05, cols=2, card_h=2.25, title_s=14, body_s=11.5, gap=0.28):
    rows = (len(cards) + cols - 1) // cols
    w = (SW - 2 * M - gap * (cols - 1)) / cols
    for i, (t, d) in enumerate(cards):
        cx = M + (i % cols) * (w + gap)
        cy = y0 + (i // cols) * (card_h + gap)
        box(slide, cx, cy, w, card_h, fill=WHITE, line=BORDER, round_=True, adj=0.05)
        box(slide, cx + 0.28, cy + 0.26, 0.34, 0.055, fill=TEAL)
        tb(slide, cx + 0.28, cy + 0.42, w - 0.56, 0.4, [{"runs": [P(t, title_s, True, False, SLATE)]}])
        tb(slide, cx + 0.28, cy + 0.9, w - 0.56, card_h - 1.05,
           [{"runs": [P(d, body_s, False, False, SLATE_6)], "line": 1.2}])


def render_title(s, sl):
    box(s, 0, 0, SW, 0.09, fill=TEAL)
    tb(s, M, 1.0, 8, 0.4, [{"runs": [P(sl["kicker"], 12, True, False, TEAL_3)]}])
    if os.path.exists(LOGO_REVERSE):
        s.shapes.add_picture(LOGO_REVERSE, Inches(M), Inches(1.45), width=Inches(5.6))
    box(s, M + 0.06, 3.5, 0.62, 0.06, fill=TEAL)
    tb(s, M, 3.74, 7.4, 0.9, [{"runs": [P(sl["sub"], 20, True, False, TEAL_3)], "line": 1.2}])
    tb(s, M, 4.75, 7.0, 1.2, [
        {"runs": [P("AI question papers · secure online exams · auto-grading · fees · attendance · results",
                    12.5, False, False, SLATE_3)], "line": 1.35, "space_after": 4},
        {"runs": [P("Six Indian languages · UUCMS-ready · ₹449 per student per year", 12.5, False, False, SLATE_4)],
         "line": 1.35},
    ])
    card_x = 8.5
    box(s, card_x, 1.15, 4.2, 5.2, fill=WHITE, round_=True, adj=0.045)
    if os.path.exists(HERO):
        hw = 3.9
        s.shapes.add_picture(HERO, Inches(card_x + (4.2 - hw) / 2), Inches(2.0), width=Inches(hw))
    if os.path.exists(LOGO_LIGHT):
        s.shapes.add_picture(LOGO_LIGHT, Inches(card_x + 0.75), Inches(5.35), width=Inches(2.7))
    tb(s, M, 6.75, 7.4, 0.5, [{"runs": [P(sl["foot"], 10.5, False, True, SLATE_4)]}])


def render_stats(s, sl):
    header(s, sl["kicker"], sl["title"], sl.get("sub"))
    stat_row(s, 2.02, sl["stats"], box_h=1.42)
    bullet_list(s, M, 3.72, SW - 2 * M - 0.2, sl["bullets"], s=12, gap=8)


def render_cards(s, sl):
    header(s, sl["kicker"], sl["title"], sl.get("sub"))
    n = len(sl["cards"])
    cols = 2 if n <= 4 else 3
    card_h = 2.2 if n <= 4 else 2.05
    cards = sl["cards"]
    bullets = sl.get("bullets") or []
    if bullets:                       # keep the cards compact so the bullets fit
        card_h = min(card_h, 1.95)
    card_grid(s, cards, y0=2.02, cols=cols, card_h=card_h,
              body_s=11.5 if n <= 4 else 11)
    if bullets:
        rows = (len(cards) + cols - 1) // cols
        y = 2.02 + rows * card_h + (rows - 1) * 0.28 + 0.28
        bullet_list(s, M, y, SW - 2 * M - 0.2, bullets, s=11, gap=6,
                    h=max(0.8, SH - y - 0.55))


def render_flow(s, sl):
    header(s, sl["kicker"], sl["title"], sl.get("sub"))
    steps = sl["steps"]
    n = len(steps)
    gap = 0.18
    w = (SW - 2 * M - gap * (n - 1)) / n
    for i, (t, d) in enumerate(steps):
        x = M + i * (w + gap)
        box(s, x, 2.02, w, 2.55, fill=WHITE, line=BORDER, round_=True, adj=0.06)
        box(s, x, 2.02, w, 0.06, fill=TEAL)
        tb(s, x + 0.2, 2.24, w - 0.4, 0.6, [{"runs": [P(t, 13, True, False, TEAL_D)]}])
        tb(s, x + 0.2, 2.82, w - 0.4, 1.6, [{"runs": [P(d, 10.5, False, False, SLATE_6)], "line": 1.2}])
    bullet_list(s, M, 4.85, SW - 2 * M - 0.2, sl["bullets"], s=11.5, gap=7)


def render_split(s, sl):
    header(s, sl["kicker"], sl["title"], sl.get("sub"))
    lw = 7.35
    rw = SW - 2 * M - lw - 0.35
    rx = M + lw + 0.35
    tb(s, M, 2.02, lw, 0.35, [{"runs": [P(sl["left_title"], 12, True, False, TEAL_D)]}])
    bullet_list(s, M, 2.45, lw, sl["left"], s=11.5, gap=9, h=4.2)
    box(s, rx, 2.02, rw, 4.6, fill=WHITE, line=BORDER, round_=True, adj=0.05)
    tb(s, rx + 0.3, 2.28, rw - 0.6, 0.35, [{"runs": [P(sl["right_title"], 12, True, False, TEAL_D)]}])
    panel_list(s, rx + 0.3, 2.72, rw - 0.6, sl["right"], s=11, gap=11, h=3.7)


def render_split2(s, sl):
    header(s, sl["kicker"], sl["title"], sl.get("sub"))
    lw = (SW - 2 * M - 0.3) / 2
    rx = M + lw + 0.3
    box(s, M, 2.02, lw, 4.6, fill=WHITE, line=BORDER, round_=True, adj=0.05)
    box(s, M, 2.02, 0.055, 4.6, fill=TEAL)
    tb(s, M + 0.3, 2.26, lw - 0.6, 0.35, [{"runs": [P(sl["left_title"], 12.5, True, False, TEAL_D)]}])
    bullet_list(s, M + 0.3, 2.72, lw - 0.6, sl["left"], s=11, gap=9, h=3.7)
    box(s, rx, 2.02, lw, 4.6, fill=DARK_BG, round_=True, adj=0.05)
    tb(s, rx + 0.3, 2.26, lw - 0.6, 0.35, [{"runs": [P(sl["right_title"], 12.5, True, False, TEAL_3)]}])
    panel_list(s, rx + 0.3, 2.72, lw - 0.6, sl["right"], s=11, gap=10,
               title_c=TEAL_3, body_c=SLATE_3, h=3.7)


def render_table(s, sl):
    header(s, sl["kicker"], sl["title"], sl.get("sub"))
    headers = sl["headers"]
    rows = sl["rows"]
    widths = sl["widths"]
    n_rows = len(rows) + 1
    tbl_h = 0.42 + 0.42 * len(rows)
    gfx = s.shapes.add_table(n_rows, len(headers), Inches(M), Inches(2.0),
                             Inches(sum(widths)), Inches(tbl_h))
    table = gfx.table
    table.first_row = False
    table.horz_banding = False
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    table.rows[0].height = Inches(0.42)
    for r in range(1, n_rows):
        table.rows[r].height = Inches(0.42)

    def set_cell(cell, text, size, bold, color, fill, align=PP_ALIGN.LEFT):
        cell.fill.solid()
        cell.fill.fore_color.rgb = C(fill)
        cell.margin_left = cell.margin_right = Inches(0.09)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = C(color)

    for i, h in enumerate(headers):
        align = PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER
        fill = DARK_BG if i == len(headers) - 1 else SLATE_2
        set_cell(table.cell(0, i), h, 11.5, True, WHITE, fill, align)
    for r, row in enumerate(rows, start=1):
        fill = TEAL_50 if r % 2 else WHITE
        for i, val in enumerate(row):
            align = PP_ALIGN.LEFT if i == 0 else PP_ALIGN.CENTER
            last = i == len(headers) - 1
            set_cell(table.cell(r, i), val, 10.5, last, TEAL_D if last else SLATE_6, fill, align)


def render_closing(s, sl):
    box(s, 0, SH - 0.09, SW, 0.09, fill=TEAL)
    box(s, 0, 0, SW, 0.03, fill=TEAL)
    if os.path.exists(LOGO_REVERSE):
        s.shapes.add_picture(LOGO_REVERSE, Inches(M), Inches(0.7), width=Inches(2.6))
    tb(s, M, 2.35, 7.6, 1.0, [{"runs": [P(sl["title"], 42, True, False, WHITE)]}])
    box(s, M + 0.06, 3.55, 0.62, 0.06, fill=TEAL)
    tb(s, M, 3.9, 7.6, 2.0, [
        {"runs": [P(sl["sub"], 15, True, False, TEAL_3)], "space_after": 12, "line": 1.3},
        {"runs": [P(sl["lines"][0], 12.5, False, False, SLATE_3)], "space_after": 8, "line": 1.3},
        {"runs": [P(sl["lines"][1], 12.5, False, False, SLATE_4)], "line": 1.3},
    ])
    if os.path.exists(HERO):
        s.shapes.add_picture(HERO, Inches(8.9), Inches(2.4), width=Inches(3.6))


RENDERERS = {
    "title": render_title,
    "stats": render_stats,
    "cards": render_cards,
    "flow": render_flow,
    "split": render_split,
    "split2": render_split2,
    "table": render_table,
    "closing": render_closing,
}


def build_pptx(path):
    prs2 = Presentation()
    prs2.slide_width = Inches(SW)
    prs2.slide_height = Inches(SH)
    globals()["prs"] = prs2
    for sl in SLIDES:
        dark = sl["kind"] in ("title", "closing")
        s = slide_bg(DARK_BG if dark else BG)
        RENDERERS[sl["kind"]](s, sl)
    prs2.save(path)
    print("saved", path, f"({len(SLIDES)} slides)")


# ═════════════════════════════════════════════════════════════════════════════
# 3. HTML RENDERER  (self-contained; arrow keys / space navigate, Ctrl+P → PDF)
# ═════════════════════════════════════════════════════════════════════════════

CSS = """
:root{--teal:#14B8A6;--teal-d:#0F766E;--teal3:#5EEAD4;--teal50:#F0FDFA;--slate:#0F172A;
--slate6:#475569;--slate5:#64748B;--slate4:#94A3B8;--slate3:#CBD5E1;--border:#E2E8F0;
--bg:#F8FAFC;--dark:#0B1220;--white:#fff}
*{box-sizing:border-box}
html,body{margin:0;height:100%;background:#0b1220;
font-family:Inter,-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,"Noto Sans",sans-serif;color:var(--slate)}
.deck{height:100vh;overflow-y:auto;scroll-snap-type:y mandatory}
.slide{width:100%;min-height:100vh;scroll-snap-align:start;display:flex;justify-content:center;
align-items:center;padding:24px}
.inner{position:relative;width:min(1280px,96vw);aspect-ratio:16/9;background:var(--bg);border-radius:14px;
overflow:hidden;box-shadow:0 24px 60px rgba(0,0,0,.45);padding:44px 52px;display:flex;flex-direction:column}
.inner.dark{background:var(--dark);color:#fff}
.kicker{color:var(--teal-d);font-size:12px;font-weight:700;letter-spacing:.12em;text-transform:uppercase}
.inner.dark .kicker{color:var(--teal3)}
.rule{width:44px;height:4px;background:var(--teal);border-radius:2px;margin-bottom:14px}
h1{font-size:34px;line-height:1.15;margin:6px 0 0;letter-spacing:-.01em}
.sub{color:var(--slate5);font-size:15px;margin-top:10px}
.inner.dark .sub{color:var(--slate3)}
.grid{display:grid;gap:16px;margin-top:22px}
.g2{grid-template-columns:1fr 1fr}
.g3{grid-template-columns:repeat(3,1fr)}
.g4{grid-template-columns:repeat(4,1fr)}
.g5{grid-template-columns:repeat(5,1fr)}
.card{background:#fff;border:1px solid var(--border);border-radius:12px;padding:16px 18px}
.card .t{font-weight:700;font-size:15px;margin-bottom:6px}
.card .d{color:var(--slate6);font-size:12.5px;line-height:1.4}
.card .bar{width:26px;height:3px;background:var(--teal);border-radius:2px;margin-bottom:10px}
.stat{background:#fff;border:1px solid var(--border);border-radius:12px;padding:14px 16px;border-left:4px solid var(--teal)}
.stat .v{color:var(--teal-d);font-size:26px;font-weight:700;line-height:1.1}
.stat .l{color:var(--slate6);font-size:11.5px;line-height:1.35;margin-top:6px}
ul.b{list-style:none;margin:18px 0 0;padding:0}
ul.b li{position:relative;padding-left:20px;margin-bottom:9px;color:var(--slate6);font-size:13px;line-height:1.4}
ul.b li:before{content:"";position:absolute;left:0;top:7px;width:7px;height:7px;border-radius:50%;background:var(--teal)}
.panel{background:#fff;border:1px solid var(--border);border-radius:12px;padding:18px}
.panel.dark{background:var(--dark);border-color:#1e293b}
.panel.dark .t{color:var(--teal3)}.panel.dark .d{color:var(--slate3)}
.panel .t{font-weight:700;color:var(--teal-d);font-size:14px;margin-bottom:4px}
.panel .d{color:var(--slate6);font-size:12.5px;line-height:1.4;margin-bottom:12px}
table{width:100%;border-collapse:collapse;margin-top:20px;font-size:12.5px}
th,td{padding:9px 12px;text-align:center;border-bottom:1px solid var(--border)}
th:first-child,td:first-child{text-align:left}
thead th{background:var(--slate);color:#fff;font-size:11.5px;letter-spacing:.03em}
thead th:last-child{background:var(--dark)}
tbody tr:nth-child(odd){background:var(--teal50)}
tbody td:last-child{color:var(--teal-d);font-weight:700}
.h22{display:grid;grid-template-columns:1fr 1fr;gap:18px;margin-top:22px}
.titleWrap{margin-top:auto;margin-bottom:auto;max-width:56%}
.big{font-size:74px;font-weight:800;letter-spacing:-.03em;line-height:1}
.footer{margin-top:auto;color:var(--slate4);font-size:11px}
.nav{position:fixed;right:18px;bottom:16px;display:flex;gap:8px;align-items:center;z-index:9;
background:rgba(11,18,32,.72);border:1px solid #1e293b;border-radius:999px;padding:6px 12px;
color:var(--slate3);font-size:12px;backdrop-filter:blur(6px)}
.nav button{background:transparent;border:0;color:var(--teal3);font-size:15px;cursor:pointer;padding:0 4px}
.hint{position:fixed;left:18px;bottom:16px;color:var(--slate4);font-size:11px;z-index:9}
@media print{
  html,body{background:#fff}
  .nav,.hint{display:none}
  .deck{overflow:visible;height:auto}
  .slide{min-height:auto;padding:0;page-break-after:always}
  .inner{box-shadow:none;border-radius:0;width:100%;height:100vh}
  @page{size:A4 landscape;margin:0}
}
"""


def esc(t):
    return html.escape(str(t))


def html_slide(sl, idx, total):
    kind = sl["kind"]
    dark = kind in ("title", "closing")
    cls = "inner dark" if dark else "inner"
    out = [f'<section class="slide" id="s{idx + 1}"><div class="{cls}">']

    if kind == "title":
        out.append('<div style="width:44px;height:4px;background:var(--teal);border-radius:2px"></div>')
        out.append(f'<div class="kicker" style="margin-top:14px">{esc(sl["kicker"])}</div>')
        out.append('<div class="titleWrap">')
        out.append(f'<div class="big">{esc(sl["title"])}</div>')
        out.append(f'<h1 style="color:var(--teal3);font-size:24px;margin-top:18px">{esc(sl["sub"])}</h1>')
        out.append('<p class="sub" style="margin-top:18px">AI question papers · secure online exams · auto-grading · '
                   'fees · attendance · results · six Indian languages · UUCMS-ready</p>')
        out.append('</div>')
        out.append(f'<div class="footer">{esc(sl["foot"])}</div>')
    elif kind == "closing":
        out.append(f'<div class="titleWrap"><h1 style="font-size:56px">{esc(sl["title"])}</h1>'
                   f'<h1 style="color:var(--teal3);font-size:20px;margin-top:20px">{esc(sl["sub"])}</h1>')
        for ln in sl["lines"]:
            out.append(f'<p class="sub" style="margin-top:16px">{esc(ln)}</p>')
        out.append('</div>')
    else:
        out.append('<div class="rule"></div>')
        out.append(f'<div class="kicker">{esc(sl["kicker"])}</div>')
        out.append(f'<h1>{esc(sl["title"])}</h1>')
        if sl.get("sub"):
            out.append(f'<div class="sub">{esc(sl["sub"])}</div>')
        if kind == "stats":
            out.append(f'<div class="grid g{len(sl["stats"])}">')
            for v, l in sl["stats"]:
                out.append(f'<div class="stat"><div class="v">{esc(v)}</div><div class="l">{esc(l)}</div></div>')
            out.append('</div>')
            out.append('<ul class="b">' + "".join(f"<li>{esc(b)}</li>" for b in sl["bullets"]) + '</ul>')
        elif kind == "cards":
            cards = sl["cards"]
            cols = 2 if len(cards) <= 4 else 3
            out.append(f'<div class="grid g{cols}">')
            for t, d in cards:
                out.append(f'<div class="card"><div class="bar"></div><div class="t">{esc(t)}</div>'
                           f'<div class="d">{esc(d)}</div></div>')
            out.append('</div>')
            if sl.get("bullets"):
                out.append('<ul class="b">' + "".join(f"<li>{esc(b)}</li>" for b in sl["bullets"]) + '</ul>')
        elif kind == "flow":
            out.append(f'<div class="grid g{len(sl["steps"])}">')
            for t, d in sl["steps"]:
                out.append(f'<div class="card"><div class="bar"></div><div class="t" style="color:var(--teal-d)">{esc(t)}</div>'
                           f'<div class="d">{esc(d)}</div></div>')
            out.append('</div>')
            out.append('<ul class="b">' + "".join(f"<li>{esc(b)}</li>" for b in sl["bullets"]) + '</ul>')
        elif kind == "table":
            out.append('<table><thead><tr>' + "".join(f'<th>{esc(h)}</th>' for h in sl["headers"]) + '</tr></thead><tbody>')
            for row in sl["rows"]:
                out.append('<tr>' + "".join(f'<td>{esc(c)}</td>' for c in row) + '</tr>')
            out.append('</tbody></table>')
        elif kind in ("split", "split2"):
            out.append('<div class="h22">')
            left = '<div><div class="t" style="font-weight:700;color:var(--teal-d);font-size:14px;margin-bottom:6px">' \
                   + esc(sl["left_title"]) + '</div><ul class="b" style="margin-top:6px">' \
                   + "".join(f'<li>{esc(b)}</li>' for b in sl["left"]) + '</ul></div>'
            out.append(left)
            pd = "panel dark" if kind == "split2" else "panel"
            right = f'<div class="{pd}"><div class="t">{esc(sl["right_title"])}</div>'
            for t, d in sl["right"]:
                right += f'<div class="t" style="font-size:13px">{esc(t)}</div><div class="d">{esc(d)}</div>'
            right += '</div>'
            out.append(right)
            out.append('</div>')
    out.append('</div></section>')
    return "\n".join(out)


def build_html(path):
    slides = "\n".join(html_slide(sl, i, len(SLIDES)) for i, sl in enumerate(SLIDES))
    doc = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Vriddhi — Investor Deck 2026</title>
<style>{CSS}</style></head>
<body>
<div class="deck" id="deck">
{slides}
</div>
<div class="hint">← → or space to move · Ctrl/⌘+P prints to PDF</div>
<div class="nav"><button onclick="go(-1)">←</button><span id="pos">1 / {len(SLIDES)}</span>
<button onclick="go(1)">→</button></div>
<script>
const slides=[...document.querySelectorAll('.slide')];let i=0;
function go(d){{i=Math.max(0,Math.min(slides.length-1,i+d));
document.getElementById('pos').textContent=(i+1)+' / '+slides.length;
slides[i].scrollIntoView({{behavior:'smooth'}});}}
document.addEventListener('keydown',e=>{{
 if(['ArrowRight','ArrowDown',' ','PageDown'].includes(e.key)){{e.preventDefault();go(1)}}
 if(['ArrowLeft','ArrowUp','PageUp'].includes(e.key)){{e.preventDefault();go(-1)}}
 if(e.key==='Home'){{i=0;slides[0].scrollIntoView()}}
 if(e.key==='End'){{i=slides.length-1;slides[i].scrollIntoView()}}}});
addEventListener('scroll',()=>{{}});
</script>
</body></html>
"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(doc)
    print("saved", path)


# ═════════════════════════════════════════════════════════════════════════════
# 4. SPEAKER NOTES
# ═════════════════════════════════════════════════════════════════════════════

def build_notes(path):
    lines = [
        "# Vriddhi — investor deck speaker notes",
        "",
        "Companion to `Vriddhi_Investor_Deck.pptx` and `investor-deck.html`.",
        "Replace every `<angle bracket>` placeholder before you send or present this deck.",
        "",
        "**Total run time:** 10–12 minutes for the deck, then let them question you for 20.",
        "**Rule:** after each slide's 'say' line, stop talking. Silence is where they lean in.",
        "",
        "---",
        "",
    ]
    for i, sl in enumerate(SLIDES, start=1):
        title = sl.get("title") or "(cover)"
        lines.append(f"## Slide {i} · {title}")
        lines.append("")
        if sl.get("sub"):
            lines.append(f"*On screen:* {sl['sub']}")
            lines.append("")
        lines.append(sl["notes"])
        lines.append("")
    lines += [
        "---",
        "",
        "## If you only get five minutes",
        "",
        "1. Slide 1 — what it is, in one sentence.",
        "2. Slide 4 — the wedge (exam paper first, campus second, learning third).",
        "3. Slide 5 — 160 functions, 257 K lines, 1,540 tests: it is built.",
        "4. Slide 11 — ₹18,123/year infrastructure and ₹23/student AI.",
        "5. Slide 12 — say the gap first: 'no paying college yet; this round fixes that.'",
        "",
        "Then ask for the three introductions. A five-minute meeting that ends with names in your",
        "inbox beats a fifty-minute meeting that ends with 'very interesting'.",
        "",
    ]
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    print("saved", path)


if __name__ == "__main__":
    build_pptx(os.path.join(HERE, "Vriddhi_Investor_Deck.pptx"))
    build_html(os.path.join(HERE, "investor-deck.html"))
    build_notes(os.path.join(HERE, "Vriddhi_Investor_Deck_Speaker_Notes.md"))
